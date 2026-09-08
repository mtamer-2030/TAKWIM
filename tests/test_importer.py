"""اختبار محرك الاستيراد متعدد الصيغ (المرحلة 3)."""

import io

import pytest

from app.importer import (
    ExtractedText,
    ImporterError,
    extract_text,
    extract_text_from_docx,
    extract_text_from_image,
    extract_text_from_pptx,
    parse_students_excel,
)


def _xlsx(rows: list[dict]) -> bytes:
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active
    headers = list(rows[0].keys())
    ws.append(headers)
    for r in rows:
        ws.append([r.get(h, "") for h in headers])
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def _docx(paras: list[str]) -> bytes:
    from docx import Document
    d = Document()
    for p in paras:
        d.add_paragraph(p)
    buf = io.BytesIO(); d.save(buf); return buf.getvalue()


def _pptx(slides: list[list[str]]) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    blank = prs.slide_layouts[6]
    for lines in slides:
        s = prs.slides.add_slide(blank)
        tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(4))
        tf = tb.text_frame
        tf.text = lines[0]
        for extra in lines[1:]:
            tf.add_paragraph().text = extra
    buf = io.BytesIO(); prs.save(buf); return buf.getvalue()


# ————— Excel لوائح —————

def test_excel_roster_arabic_headers():
    data = _xlsx([
        {"الاسم والنسب": "الإدريسي صفوان", "رقم مسار": "D161055238", "القسم": "TC1"},
        {"الاسم والنسب": "الورد علياء", "رقم مسار": "D167089903", "القسم": "TC1"},
    ])
    r = parse_students_excel(data)
    assert r.ok, r.errors
    assert len(r.rows) == 2
    assert r.rows[0].full_name == "الإدريسي صفوان"
    assert r.rows[0].massar_code == "D161055238"
    assert r.rows[0].group_name == "TC1"


def test_excel_english_headers_and_blank_rows():
    data = _xlsx([
        {"full_name": "Ali", "massar": "X1", "group": "1BAC2"},
        {"full_name": "", "massar": "", "group": ""},   # صفّ فارغ يُتجاهَل
    ])
    r = parse_students_excel(data)
    assert r.ok and len(r.rows) == 1 and r.rows[0].massar_code == "X1"


def test_excel_missing_name_column_rejected():
    data = _xlsx([{"مسار": "D1", "القسم": "TC1"}])
    r = parse_students_excel(data)
    assert not r.ok and any("الاسم الكامل" in e for e in r.errors)


def test_excel_duplicate_massar_flagged():
    data = _xlsx([
        {"الاسم": "أ", "مسار": "D1"},
        {"الاسم": "ب", "مسار": "D1"},
    ])
    r = parse_students_excel(data)
    assert not r.ok and any("مكرّر" in e and "السطر 3" in e for e in r.errors)


# ————— استخراج النصوص —————

def test_extract_docx():
    data = _docx(["الفقرة الأولى", "الفقرة الثانية"])
    out = extract_text_from_docx(data)
    assert "الفقرة الأولى" in out and "الفقرة الثانية" in out


def test_extract_pptx():
    data = _pptx([["عنوان الشريحة", "سطر ثانٍ"]])
    out = extract_text_from_pptx(data)
    assert "عنوان الشريحة" in out and "شريحة 1" in out


def test_dispatch_by_extension():
    data = _docx(["نص تجريبي"])
    res = extract_text("lesson.docx", data)
    assert isinstance(res, ExtractedText) and res.kind == "docx"
    assert "نص تجريبي" in res.text


def test_unsupported_extension_rejected():
    with pytest.raises(ImporterError):
        extract_text("file.txt", b"hello")


def test_corrupt_docx_reported():
    with pytest.raises(ImporterError):
        extract_text_from_docx(b"not a docx")


def test_image_ocr_without_tesseract_gives_clear_error():
    # صورة PNG صغيرة صالحة؛ في غياب ثنائي Tesseract يجب أن يُرفع خطأ واضح لا انهيار.
    from PIL import Image
    buf = io.BytesIO(); Image.new("RGB", (20, 20), "white").save(buf, format="PNG")
    try:
        extract_text_from_image(buf.getvalue())
    except ImporterError as e:
        assert "Tesseract" in str(e) or "تعذّر" in str(e)
