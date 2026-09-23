"""تصدير تقارير الأستاذ إلى Word (.docx) ونصّ (.txt) — جماعيّ (فوج) وفرديّ (تلميذ).

تقارير للأستاذ وحده (نقط صريحة، ق-٤). تُبنى من مخرجات services.gradebook دون تغيير
مخطّط. يُضبط الاتجاه RTL في الفقرات والجداول لتظهر العربية سليمةً في Word.
"""

from __future__ import annotations

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Pt


def _rtl_para(p):
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pPr = p._p.get_or_add_pPr()
    pPr.append(OxmlElement("w:bidi"))


def _rtl_table(table):
    tblPr = table._tbl.tblPr
    tblPr.append(OxmlElement("w:bidiVisual"))


def _pct(v) -> str:
    return f"{v}٪" if v is not None else "—"


def _new_doc(title: str) -> Document:
    doc = Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(11)
    h = doc.add_heading(title, level=1)
    _rtl_para(h)
    return doc


# ═══════════════ تقرير التلميذ ═══════════════


def student_report_docx(data: dict) -> bytes:
    st = data["student"]
    doc = _new_doc(f"التقرير الفرديّ — {st.full_name}")
    meta = doc.add_paragraph()
    _rtl_para(meta)
    trend = data.get("trend")
    trend_txt = (("+" if trend >= 0 else "") + f"{trend} نقطة") if trend is not None else "—"
    meta.add_run(
        f"الفوج: {st.group_name or '—'}  |  المعدّل العامّ: {_pct(data.get('overall'))}"
        f"  |  التطوّر (أوّل→آخر): {trend_txt}")

    # جدول التقويمات
    ph = doc.add_paragraph(); _rtl_para(ph); ph.add_run("النسب عبر التقويمات:").bold = True
    evals = data.get("evaluations") or []
    if evals:
        t = doc.add_table(rows=1, cols=4); t.style = "Table Grid"; _rtl_table(t)
        for i, head in enumerate(["التاريخ", "التقويم", "النوع", "النسبة"]):
            t.rows[0].cells[i].paragraphs[0].add_run(head).bold = True
            _rtl_para(t.rows[0].cells[i].paragraphs[0])
        for e in evals:
            c = t.add_row().cells
            for i, val in enumerate([e.get("date", "—"), e.get("title", ""),
                                     e.get("kind", ""), _pct(e.get("pct"))]):
                c[i].paragraphs[0].add_run(str(val)); _rtl_para(c[i].paragraphs[0])
    else:
        p = doc.add_paragraph("لا تقويمات مصادَق عليها بعد."); _rtl_para(p)

    # جدول المهارات
    skills = data.get("skills") or {}
    if skills:
        ps = doc.add_paragraph(); _rtl_para(ps); ps.add_run("ملمح المهارات:").bold = True
        t = doc.add_table(rows=1, cols=2); t.style = "Table Grid"; _rtl_table(t)
        for i, head in enumerate(["المهارة", "التمكّن ٪"]):
            t.rows[0].cells[i].paragraphs[0].add_run(head).bold = True
            _rtl_para(t.rows[0].cells[i].paragraphs[0])
        for name, sk in skills.items():
            avg = sk.get("avg")
            c = t.add_row().cells
            c[0].paragraphs[0].add_run(str(name)); _rtl_para(c[0].paragraphs[0])
            c[1].paragraphs[0].add_run(_pct(round(avg * 100, 1) if avg is not None else None))
            _rtl_para(c[1].paragraphs[0])

    buf = io.BytesIO(); doc.save(buf); return buf.getvalue()


def student_report_txt(data: dict) -> str:
    st = data["student"]
    lines = [f"التقرير الفرديّ — {st.full_name}",
             f"الفوج: {st.group_name or '—'}",
             f"المعدّل العامّ: {_pct(data.get('overall'))}",
             f"التطوّر: {data.get('trend') if data.get('trend') is not None else '—'}",
             "", "النسب عبر التقويمات:"]
    for e in (data.get("evaluations") or []):
        lines.append(f"  - {e.get('date','—')} | {e.get('title','')} | {e.get('kind','')} | {_pct(e.get('pct'))}")
    if not (data.get("evaluations")):
        lines.append("  (لا تقويمات مصادَق عليها بعد)")
    lines += ["", "ملمح المهارات:"]
    for name, sk in (data.get("skills") or {}).items():
        avg = sk.get("avg")
        lines.append(f"  - {name}: {_pct(round(avg*100,1) if avg is not None else None)}")
    return "\n".join(lines)


# ═══════════════ تقرير الفوج ═══════════════


def class_report_docx(data: dict) -> bytes:
    group = data.get("group_name", "")
    doc = _new_doc(f"التقرير الجماعيّ — الفوج {group}")
    meta = doc.add_paragraph(); _rtl_para(meta)
    meta.add_run(f"معدّل القسم العامّ: {_pct(data.get('class_overall'))}").bold = True

    evals = data.get("evaluations") or []
    students = data.get("students") or []
    cols = ["التلميذ"] + [f"{e['title']} ({e['date']})" for e in evals] + ["المعدّل ٪"]
    t = doc.add_table(rows=1, cols=len(cols)); t.style = "Table Grid"; _rtl_table(t)
    for i, head in enumerate(cols):
        t.rows[0].cells[i].paragraphs[0].add_run(head).bold = True
        _rtl_para(t.rows[0].cells[i].paragraphs[0])
    for srow in students:
        c = t.add_row().cells
        c[0].paragraphs[0].add_run(srow["student"].full_name); _rtl_para(c[0].paragraphs[0])
        for j, e in enumerate(evals):
            v = srow["per"].get(e["quiz_id"])
            c[j + 1].paragraphs[0].add_run(_pct(v)); _rtl_para(c[j + 1].paragraphs[0])
        c[-1].paragraphs[0].add_run(_pct(srow.get("overall"))); _rtl_para(c[-1].paragraphs[0])
    # صفّ معدّل القسم
    c = t.add_row().cells
    c[0].paragraphs[0].add_run("معدّل القسم").bold = True; _rtl_para(c[0].paragraphs[0])
    for j, e in enumerate(evals):
        c[j + 1].paragraphs[0].add_run(_pct(e.get("class_avg"))); _rtl_para(c[j + 1].paragraphs[0])
    c[-1].paragraphs[0].add_run(_pct(data.get("class_overall"))); _rtl_para(c[-1].paragraphs[0])

    buf = io.BytesIO(); doc.save(buf); return buf.getvalue()


def _bold_para(doc, text: str):
    p = doc.add_paragraph(); _rtl_para(p); p.add_run(text).bold = True
    return p


def ai_reports_docx(data: dict) -> bytes:
    """تقرير التدخّل الكامل ذو القيمة البيداغوجيّة في Word: تقرير جماعيّ (معدّل، تمكّن
    الكفايات بتقديرٍ نوعيّ، توزيع المستويات، القصور المهيمن، توصيات دعم) ثمّ تقريرٌ فرديّ
    لكلّ تلميذ (معدّله ووصفه، قوّته وقصوره، تمكّن كفاياته، توصيات) — كلّه بلا Ollama."""
    group = data.get("group", "")
    cr = data.get("class") or {}
    cp = cr.get("peda") or {}
    doc = _new_doc(f"تقرير التدخّل العلاجيّ — الفوج {group}")

    _bold_para(doc, "أوّلاً: التقرير الجماعيّ للقسم")
    if cr.get("narrative"):
        n = doc.add_paragraph(str(cr["narrative"])); _rtl_para(n)
    meta = doc.add_paragraph(); _rtl_para(meta)
    meta.add_run(
        f"المعدّل العامّ للقسم: {_pct(cp.get('overall_pct'))}  |  "
        f"تلاميذ لهم بيانات: {cp.get('students_with_data', 0)}/{cp.get('student_count', 0)}  |  "
        f"أقوى كفاية: {cp.get('strongest') or '—'}  |  أضعف كفاية: {cp.get('weakest') or '—'}")
    if cp.get("dominant_deficit"):
        d = doc.add_paragraph(); _rtl_para(d)
        d.add_run(f"القصور المنهجيّ المهيمن: {cp['dominant_deficit']} "
                  f"(أضعف كفايةٍ لدى {_pct(cp.get('dominant_share_pct'))} من التلاميذ)")

    _bold_para(doc, "تمكّن القسم من الكفايات")
    t = doc.add_table(rows=1, cols=3); t.style = "Table Grid"; _rtl_table(t)
    for i, head in enumerate(["الكفاية", "النسبة", "التقدير"]):
        t.rows[0].cells[i].paragraphs[0].add_run(head).bold = True
        _rtl_para(t.rows[0].cells[i].paragraphs[0])
    for d in (cp.get("detail") or []):
        c = t.add_row().cells
        for i, val in enumerate([d["name"], _pct(d.get("pct")), d.get("band", "—")]):
            c[i].paragraphs[0].add_run(str(val)); _rtl_para(c[i].paragraphs[0])

    dist = cp.get("distribution") or {}
    _bold_para(doc, "توزيع مستويات التلاميذ")
    p = doc.add_paragraph("  ·  ".join(f"{k}: {v}" for k, v in dist.items())); _rtl_para(p)

    if cr.get("hardest"):
        _bold_para(doc, "أصعب الأسئلة على القسم (الأخطاء الشائعة)")
        for q in cr["hardest"]:
            p = doc.add_paragraph(f"• [{q.get('success')}٪] {q.get('prompt','')}")
            _rtl_para(p)

    if cr.get("progress") and len(cr["progress"]) > 1:
        _bold_para(doc, "تطوّر القسم عبر الزمن (ذاكرةٌ تراكميّة)")
        for pr in cr["progress"]:
            p = doc.add_paragraph(
                f"• {pr.get('label','')} ({pr.get('date','')}): "
                f"{pr.get('pct')}٪ — تراكميّاً {pr.get('cumulative')}٪")
            _rtl_para(p)

    _bold_para(doc, "توصيات الدعم البيداغوجيّ")
    for r in (cp.get("recommendations") or []):
        p = doc.add_paragraph(); _rtl_para(p)
        p.add_run(f"• {r['skill']}: ").bold = True
        p.add_run(r.get("advice", ""))
    if cr.get("plan"):
        _bold_para(doc, "خطّة الدعم (توليد ذكيّ)")
        p = doc.add_paragraph(str(cr["plan"])); _rtl_para(p)

    doc.add_page_break()
    _bold_para(doc, "ثانياً: التقارير الفردية")
    for stu in (data.get("students") or []):
        sp = stu.get("peda") or {}
        h = doc.add_heading(stu["student"].full_name, level=2); _rtl_para(h)
        if not stu.get("has_data"):
            p = doc.add_paragraph("لا بيانات مصادَقٌ عليها بعد لهذا التلميذ."); _rtl_para(p)
            continue
        if stu.get("narrative"):
            nn = doc.add_paragraph(str(stu["narrative"])); _rtl_para(nn)
        m = doc.add_paragraph(); _rtl_para(m)
        m.add_run(f"المعدّل: {_pct(sp.get('overall_pct'))} ({sp.get('overall_band','—')})  |  "
                  f"القوّة: {sp.get('strongest') or '—'}  |  القصور: {sp.get('weakest') or '—'}")
        for d in (sp.get("detail") or []):
            p = doc.add_paragraph(f"• {d['name']}: {_pct(d.get('pct'))} — {d.get('band','—')}")
            _rtl_para(p)
        for r in (sp.get("recommendations") or []):
            p = doc.add_paragraph(); _rtl_para(p)
            p.add_run(f"توصية — {r['skill']}: ").bold = True
            p.add_run(r.get("advice", ""))
        if stu.get("plan"):
            p = doc.add_paragraph(); _rtl_para(p)
            p.add_run("خطّة ذكيّة: ").bold = True
            p.add_run(str(stu["plan"]))

    buf = io.BytesIO(); doc.save(buf); return buf.getvalue()


def ai_reports_txt(data: dict) -> str:
    group = data.get("group", "")
    cr = data.get("class") or {}
    cp = cr.get("peda") or {}
    lines = [f"تقرير التدخّل العلاجيّ — الفوج {group}", "",
             "═══ أوّلاً: التقرير الجماعيّ ═══"]
    if cr.get("narrative"):
        lines += [str(cr["narrative"]), ""]
    lines += [f"المعدّل العامّ للقسم: {_pct(cp.get('overall_pct'))}",
             f"تلاميذ لهم بيانات: {cp.get('students_with_data',0)}/{cp.get('student_count',0)}",
             f"أقوى كفاية: {cp.get('strongest') or '—'} | أضعف كفاية: {cp.get('weakest') or '—'}"]
    if cp.get("dominant_deficit"):
        lines.append(f"القصور المهيمن: {cp['dominant_deficit']} "
                     f"(لدى {_pct(cp.get('dominant_share_pct'))} من التلاميذ)")
    lines.append("\nتمكّن الكفايات:")
    for d in (cp.get("detail") or []):
        lines.append(f"  • {d['name']}: {_pct(d.get('pct'))} — {d.get('band','—')}")
    dist = cp.get("distribution") or {}
    lines.append("توزيع المستويات: " + " · ".join(f"{k}: {v}" for k, v in dist.items()))
    if cr.get("hardest"):
        lines.append("\nأصعب الأسئلة (الأخطاء الشائعة):")
        for q in cr["hardest"]:
            lines.append(f"  • [{q.get('success')}٪] {q.get('prompt','')}")
    if cr.get("progress") and len(cr["progress"]) > 1:
        lines.append("\nتطوّر القسم عبر الزمن (ذاكرةٌ تراكميّة):")
        for pr in cr["progress"]:
            lines.append(f"  • {pr.get('label','')} ({pr.get('date','')}): "
                         f"{pr.get('pct')}٪ — تراكميّاً {pr.get('cumulative')}٪")
    lines.append("\nتوصيات الدعم:")
    for r in (cp.get("recommendations") or []):
        lines.append(f"  • {r['skill']}: {r.get('advice','')}")
    if cr.get("plan"):
        lines += ["\nخطّة الدعم (ذكيّة):", str(cr["plan"])]

    lines += ["", "═══ ثانياً: التقارير الفردية ═══"]
    for stu in (data.get("students") or []):
        sp = stu.get("peda") or {}
        lines.append(f"\n— {stu['student'].full_name} —")
        if not stu.get("has_data"):
            lines.append("  لا بيانات مصادَقٌ عليها بعد.")
            continue
        lines.append(f"  المعدّل: {_pct(sp.get('overall_pct'))} ({sp.get('overall_band','—')}) | "
                     f"القوّة: {sp.get('strongest') or '—'} | القصور: {sp.get('weakest') or '—'}")
        for d in (sp.get("detail") or []):
            lines.append(f"    • {d['name']}: {_pct(d.get('pct'))} — {d.get('band','—')}")
        for r in (sp.get("recommendations") or []):
            lines.append(f"    توصية — {r['skill']}: {r.get('advice','')}")
        if stu.get("plan"):
            lines.append(f"    خطّة ذكيّة: {stu['plan']}")
    return "\n".join(lines)


def class_report_txt(data: dict) -> str:
    group = data.get("group_name", "")
    evals = data.get("evaluations") or []
    lines = [f"التقرير الجماعيّ — الفوج {group}",
             f"معدّل القسم العامّ: {_pct(data.get('class_overall'))}", ""]
    heads = ["التلميذ"] + [f"{e['title']} ({e['date']})" for e in evals] + ["المعدّل"]
    lines.append(" | ".join(heads))
    for srow in (data.get("students") or []):
        row = [srow["student"].full_name]
        row += [_pct(srow["per"].get(e["quiz_id"])) for e in evals]
        row.append(_pct(srow.get("overall")))
        lines.append(" | ".join(row))
    avg = ["معدّل القسم"] + [_pct(e.get("class_avg")) for e in evals] + [_pct(data.get("class_overall"))]
    lines.append(" | ".join(avg))
    return "\n".join(lines)
